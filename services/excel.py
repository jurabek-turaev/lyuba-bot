import calendar
import os
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

import database as db
from services.worktime import shift_hours

TABLES_DIR = "data/tables"

# Columns added after "Umumiy soat". They stay empty for employees
# whose schedule has not been set yet.
PAY_HEADER = [
    "Ish grafigi",
    "Kunlik norma (soat)",
    "Oylik norma (soat)",
    "Farq (soat)",
    "Soatlik stavka (so'm)",
    "Oylik (so'm)",
]


def _pay_columns(settings, worked_days: list) -> list:
    """
    Build the schedule / norm / overtime / rate / salary cells for one employee.

    settings is (shift_start, shift_end, hourly_rate) or None.
    worked_days is the list of worked hours, one value per day the employee worked.

    Every worked day is split into two parts:
    - norm part     = min(hours worked, daily norm)
    - overtime part = max(0, hours worked - daily norm)

    So a day below the norm adds exactly the hours worked to the monthly norm
    (8 h on a 10 h norm gives 8), and monthly norm + overtime = total hours.

    - Monthly norm = sum of the norm parts.
    - Overtime ("Farq") = sum of the overtime parts, never negative.
    - Salary = monthly norm * hourly rate. Overtime is NOT paid here:
      the admin sets its price and calculates it by hand.
    """
    if settings is None:
        return [None] * len(PAY_HEADER)

    shift_start, shift_end, hourly_rate = settings
    daily_norm = shift_hours(shift_start, shift_end)
    month_norm = round(sum(min(hours, daily_norm) for hours in worked_days), 2)
    overtime = round(sum(max(0.0, hours - daily_norm) for hours in worked_days), 2)
    salary = round(month_norm * hourly_rate) if hourly_rate else None

    return [f"{shift_start}–{shift_end}", daily_norm, month_norm, overtime, hourly_rate, salary]


async def generate_monthly_tabel(year: int, month: int) -> str:
    """
    Build one tabel for the given month:
    Xodim F.I.Sh. | Username | Lavozim | Do'kon/Guruh | <every date> |
    Umumiy ish kuni | Umumiy soat | <schedule, norm, overtime, rate, salary>

    All departments (Savdo, Ishlab chiqarish, Haydovchi, I.T.R) are in one sheet,
    told apart by the "Lavozim" column.

    Returns the path of the created .xlsx file.
    """
    rows = await db.get_month_attendance(year, month)
    pay_settings = await db.get_pay_settings()

    days_in_month = calendar.monthrange(year, month)[1]
    date_columns = [date(year, month, d) for d in range(1, days_in_month + 1)]
    date_labels = [d.strftime("%d.%m.%Y") for d in date_columns]

    # user_id bo'yicha guruhlash
    employees = {}
    for user_id, full_name, username, department, branch, work_date, work_hours in rows:
        if user_id not in employees:
            employees[user_id] = {
                "name": full_name,
                "username": username if username else "Mavjud emas",
                "department": department if department else "Aniqlanmagan",
                "branch": branch,
                "hours_by_date": {},
            }

        d = datetime.strptime(work_date, "%Y-%m-%d").strftime("%d.%m.%Y")
        current = employees[user_id]["hours_by_date"].get(d, 0.0)
        employees[user_id]["hours_by_date"][d] = current + (work_hours or 0.0)

    wb = Workbook()
    ws = wb.active
    ws.title = f"{month:02d}.{year}"

    header = (
        ["Xodim F.I.Sh.", "Username", "Lavozim", "Do'kon/Guruh"]
        + date_labels
        + ["Umumiy ish kuni", "Umumiy soat"]
        + PAY_HEADER
    )
    ws.append(header)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for user_id, emp in sorted(employees.items(), key=lambda item: item[1]["name"]):
        row = [emp["name"], emp["username"], emp["department"], emp["branch"]]

        worked_days = []  # rounded hours of every day the employee worked
        total_hours = 0.0

        for label in date_labels:
            hours = emp["hours_by_date"].get(label)
            if hours:
                rounded = round(hours, 2)
                row.append(rounded)
                worked_days.append(rounded)
                total_hours += rounded
            else:
                row.append(None)

        total_hours = round(total_hours, 2)
        row.append(len(worked_days))
        row.append(total_hours)
        row.extend(_pay_columns(pay_settings.get(user_id), worked_days))
        ws.append(row)

    for i in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14

    # Thousands separators for the money columns
    for title in ("Soatlik stavka (so'm)", "Oylik (so'm)"):
        column = header.index(title) + 1
        for cells in ws.iter_rows(min_row=2, min_col=column, max_col=column):
            cells[0].number_format = "#,##0"

    os.makedirs(TABLES_DIR, exist_ok=True)
    file_path = os.path.join(TABLES_DIR, f"Tabel_{month:02d}-{year}.xlsx")
    wb.save(file_path)

    return file_path